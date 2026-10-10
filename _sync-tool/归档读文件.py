"""归档读文件助手 —— 把一个文件的内容读出来，用于"读内容 → 归类 → 写说明"。

用法：
    python 归档读文件.py <文件路径>

Windows: C:\\Users\\<你>\\.dsh\\dsh-runtimes\\dsh-primary-runtime\\dependencies\\python\\python.exe
macOS  : ~/DSH-Skills/_sync-tool/归档读文件.py  （用系统 python3 即可）

支持 .pptx / .docx / .xlsx：打印页数/段落/表格、每页或每段的文字、主题配色与字体。
图片（.png/.jpg）不在这里处理 —— 直接看图更准。
"""
import os
import re
import sys

# 保证中文输出不因控制台编码报错或变乱码
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

if len(sys.argv) < 2:
    print(__doc__)
    sys.exit(1)

path = sys.argv[1]
ext = os.path.splitext(path)[1].lower()
print(f"文件: {os.path.basename(path)}  ({round(os.path.getsize(path) / 1024, 1)} KB)")


def need(module, hint):
    """检查依赖库。缺了就给一句能照着做的提示，别把 traceback 丢给用户。"""
    try:
        __import__(module)
    except ImportError:
        print()
        print(f"[缺少依赖] 这台机器没装 {module}。两种办法：")
        print(f"  1) 装一下：{hint}")
        print("     （Mac/Linux 一次性装全：pip3 install --user python-pptx python-docx openpyxl）")
        print("  2) 或者改用技能来读：office-pptx / office-docx / office-xlsx（它们自带 Python 环境）")
        print()
        sys.exit(2)


if ext == ".pptx":
    need("pptx", "pip install python-pptx")
    from pptx import Presentation
    p = Presentation(path)
    ratio = p.slide_width / p.slide_height
    shape = "16:9" if abs(ratio - 16 / 9) < 0.05 else ("4:3" if abs(ratio - 4 / 3) < 0.05 else f"{round(ratio, 2)}:1")
    print(f"类型: PPTX   页数: {len(p.slides)}   画布: {round(p.slide_width / 914400, 2)}x{round(p.slide_height / 914400, 2)} 英寸 ({shape})")
    for i, s in enumerate(p.slides, 1):
        texts = []
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                texts.append(sh.text_frame.text.strip().replace("\n", " / ")[:50])
        pics = sum(1 for sh in s.shapes if sh.shape_type == 13)
        extra = f"  [图片x{pics}]" if pics else ""
        print(f"  第{i:>3}页: {' | '.join(texts[:3]) if texts else '(纯图页)'}{extra}")
    try:
        from pptx.opc.constants import RELATIONSHIP_TYPE as RT
        theme = p.slide_masters[0].part.part_related_by(RT.THEME)
        xml = theme.blob.decode("utf-8", "ignore")
        colors = list(dict.fromkeys(re.findall(r'<a:srgbClr val="([0-9A-Fa-f]{6})"', xml)))[:6]
        fonts = [f for f in dict.fromkeys(re.findall(r'typeface="([^"]+)"', xml)) if not f.startswith("+")][:6]
        if colors:
            print("主题色:", ", ".join("#" + c for c in colors))
        if fonts:
            print("字体  :", ", ".join(fonts))
    except Exception as e:
        print("主题  : 读不到（", e, "）")

elif ext == ".docx":
    need("docx", "pip install python-docx")
    from docx import Document
    d = Document(path)
    paras = [x.text.strip() for x in d.paragraphs if x.text.strip()]
    print(f"类型: DOCX   段落: {len(paras)}   表格: {len(d.tables)}")
    for t in paras[:15]:
        print("   ", t[:70])
    if len(paras) > 15:
        print(f"    …（还有 {len(paras) - 15} 段）")
    for i, tb in enumerate(d.tables[:2], 1):
        rows = len(tb.rows)
        head = " | ".join(c.text.strip()[:12] for c in tb.rows[0].cells) if rows else ""
        print(f"    表{i}: {rows} 行    首行: {head}")

elif ext == ".xlsx":
    need("openpyxl", "pip install openpyxl")
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    print(f"类型: XLSX   工作表: {wb.sheetnames}")
    for ws in wb.worksheets[:3]:
        print(f"  [{ws.title}] {ws.max_row} 行 x {ws.max_column} 列")
        for row in ws.iter_rows(min_row=1, max_row=min(5, ws.max_row), values_only=True):
            print("    ", " | ".join("" if c is None else str(c)[:18] for c in row))
else:
    print("类型: 其他（图片、PDF、纯文本等）—— 图片直接看图，PDF/文本用阅读工具打开")
